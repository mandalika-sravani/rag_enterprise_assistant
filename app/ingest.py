"""Read documents into LangChain Documents, split them with LangChain text splitters, attach access metadata.

Every chunk is a LangChain `Document`:
    page_content = "<title> > <section>\\n<text>"   (heading included so short chunks match their topic)
    metadata     = {id, file, title, section, page, department, access, text}
"""

import re
from pathlib import Path

import yaml
from bs4 import BeautifulSoup
from langchain_core.documents import Document
from langchain_text_splitters import HTMLHeaderTextSplitter, MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

SUPPORTED = {".md", ".txt", ".html", ".htm", ".pdf", ".docx"}
HEADERS = ["h1", "h2", "h3"]


def _section(metadata: dict) -> str:
    """The deepest heading below the document title, e.g. "2. Annual Leave"."""
    return next((metadata[h] for h in ("h3", "h2") if metadata.get(h)), "")


def load_markdown(path: Path) -> tuple[str, list[Document]]:
    splitter = MarkdownHeaderTextSplitter([("#", "h1"), ("##", "h2"), ("###", "h3")])
    parts = splitter.split_text(path.read_text(encoding="utf-8"))
    title = next((p.metadata["h1"] for p in parts if p.metadata.get("h1")), path.stem)
    return title, [Document(p.page_content, metadata={"section": _section(p.metadata)}) for p in parts]


def load_html(path: Path) -> tuple[str, list[Document]]:
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else path.stem
    for tag in soup.find_all(style=re.compile(r"display\s*:\s*none", re.I)):
        tag.decompose()  # hidden text is a common place to hide prompt injections
    for row in soup.find_all("tr"):  # keep each table row on one line: "L3 | Senior Engineer | 20 to 32 LPA"
        p = soup.new_tag("p")
        p.string = " | ".join(cell.get_text(" ", strip=True) for cell in row.find_all(["td", "th"]))
        row.replace_with(p)
    splitter = HTMLHeaderTextSplitter([(h, h) for h in HEADERS])
    parts = splitter.split_text(str(soup))
    headings = {v for p in parts for v in p.metadata.values()}
    return title, [Document(p.page_content, metadata={"section": _section(p.metadata)})
                   for p in parts if p.page_content.strip() not in headings]  # drop heading-only fragments


def load_pdf(path: Path) -> tuple[str, list[Document]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = (reader.metadata.title if reader.metadata else None) or path.stem
    docs, section = [], ""
    for page_no, page in enumerate(reader.pages, start=1):
        # PDFs have no heading markup, so split each page at numbered headings ("5. Incident Reporting").
        for block in re.split(r"\n(?=\d+\.\s+[A-Z])", page.extract_text() or ""):
            lines = block.strip().splitlines()
            if lines and re.match(r"^\d+\.\s+[A-Z].{2,60}$", lines[0]):
                section, lines = lines[0], lines[1:]
            text = "\n".join(l for l in lines if l.strip() != title)
            docs.append(Document(text, metadata={"section": section, "page": page_no}))
    return title, docs


def load_docx(path: Path) -> tuple[str, list[Document]]:
    # LangChain's DOCX loaders either drop headings (Docx2txtLoader) or need the heavy `unstructured`
    # package, so this one reader uses python-docx directly to keep headings and tables in order.
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    title = d.core_properties.title or path.stem
    docs, section, lines = [], "", []
    for element in d.element.body.iterchildren():
        if element.tag.endswith("}p"):
            para = Paragraph(element, d)
            style = para.style.name if para.style is not None else ""
            if style.startswith("Heading"):
                docs.append(Document("\n".join(lines), metadata={"section": section}))
                section, lines = para.text.strip(), []
            elif style != "Title" and para.text.strip():
                lines.append(para.text.strip())
        elif element.tag.endswith("}tbl"):
            lines += [" | ".join(c.text.strip() for c in row.cells) for row in Table(element, d).rows]
    docs.append(Document("\n".join(lines), metadata={"section": section}))
    return title, docs


LOADERS = {".md": load_markdown, ".txt": load_markdown, ".html": load_html, ".htm": load_html,
           ".pdf": load_pdf, ".docx": load_docx}


def load_folder(folder: Path, size: int = 900, overlap: int = 150) -> tuple[list[Document], list[dict]]:
    """Return chunk Documents plus a per-file report (indexed / failed / skipped)."""
    manifest_path = folder / "manifest.yaml"
    manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    defaults = manifest.get("defaults") or {"department": "unassigned", "access_level": "restricted"}
    listed = manifest.get("documents") or {}
    # LangChain splitter: tries paragraph breaks, then lines, then sentences, then words.
    splitter = RecursiveCharacterTextSplitter(chunk_size=size, chunk_overlap=overlap,
                                              separators=["\n\n", "\n", ". ", " ", ""])
    chunks, report = [], []
    for path in sorted(folder.iterdir()):
        if path.name == "manifest.yaml" or not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED:
            report.append({"file": path.name, "status": "skipped", "error": "unsupported type"})
            continue
        meta = {**defaults, **listed.get(path.name, {})}  # unlisted files get the deny-by-default values
        try:
            if path.stat().st_size == 0:
                raise ValueError("file is empty")
            title, sections = LOADERS[path.suffix.lower()](path)
            title = meta.get("title") or title
            count = 0
            for piece in splitter.split_documents([s for s in sections if s.page_content.strip()]):
                section = piece.metadata.get("section", "")
                chunks.append(Document(
                    page_content=f"{title} > {section}\n{piece.page_content}" if section else f"{title}\n{piece.page_content}",
                    metadata={"id": f"{path.name}#{count}", "file": path.name, "title": title, "section": section,
                              "page": piece.metadata.get("page"), "department": meta["department"],
                              "access": meta["access_level"], "text": piece.page_content}))
                count += 1
            if count == 0:
                raise ValueError("no text found")
            report.append({"file": path.name, "status": "indexed", "chunks": count})
        except Exception as exc:  # one bad file must not stop the others
            report.append({"file": path.name, "status": "failed", "error": f"{type(exc).__name__}: {exc}"[:200]})
    return chunks, report
